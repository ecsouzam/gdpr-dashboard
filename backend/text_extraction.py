"""
Text extraction helpers for uploaded architecture documentation
(.pdf, .docx, .txt).
"""
import io
import logging

from docx import Document
from pypdf import PdfReader

logger = logging.getLogger("gdpr_dashboard.extraction")

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class UnsupportedFileTypeError(Exception):
    pass


class EmptyDocumentError(Exception):
    pass


def extract_text(filename: str, file_bytes: bytes) -> str:
    """
    Extract raw text from an uploaded file based on its extension.
    Raises UnsupportedFileTypeError or EmptyDocumentError on failure.
    """
    lower_name = filename.lower()

    if lower_name.endswith(".pdf"):
        text = _extract_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        text = _extract_docx(file_bytes)
    elif lower_name.endswith(".txt"):
        text = _extract_txt(file_bytes)
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file type for '{filename}'. Please upload a .pdf, .docx, or .txt file."
        )

    text = text.strip()
    if not text:
        raise EmptyDocumentError(
            f"No extractable text was found in '{filename}'. The file may be empty, "
            f"image-only, or corrupted."
        )

    logger.info("Extracted %d characters of text from '%s'", len(text), filename)
    return text


def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages_text = []
    for page in reader.pages:
        pages_text.append(page.extract_text() or "")
    return "\n".join(pages_text)


def _extract_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    table_cells = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    table_cells.append(cell.text)

    return "\n".join(paragraphs + table_cells)


def _extract_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")
